from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\Users\Salam\Documents\ECGdetector\output\patient_specific_conduction_repolarization_review.docx")

BLUE = "1F4E79"
DARK_BLUE = "17365D"
PALE_BLUE = "EAF2F8"
PALE_GREEN = "E9F4EC"
PALE_AMBER = "FFF4D6"
PALE_RED = "FCE8E6"
LIGHT_GRAY = "F4F6F9"
MID_GRAY = "D9E1E8"
DARK_GRAY = "3F4A54"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, twips: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(twips))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, keep=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    node = p_pr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        p_pr.append(node)
    node.set(qn("w:val"), "1" if keep else "0")


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def set_table_borders(table, color=MID_GRAY, size=4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_table_indent(table, twips=120) -> None:
    tbl_pr = table._tbl.tblPr
    indent = tbl_pr.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:w"), str(twips))
    indent.set(qn("w:type"), "dxa")


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def add_hyperlink(paragraph, text: str, url: str, color=BLUE, underline=True):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def set_section_geometry(section, landscape=False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("1F252B")
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.333

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(16)

    subtitle = styles.add_style("Report Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(15)
    subtitle.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.line_spacing = 1.15

    small = styles.add_style("Small Body", WD_STYLE_TYPE.PARAGRAPH)
    small.font.name = "Calibri"
    small.font.size = Pt(9)
    small.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.05

    table_text = styles.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    table_text.font.name = "Calibri"
    table_text.font.size = Pt(8.5)
    table_text.paragraph_format.space_after = Pt(2)
    table_text.paragraph_format.line_spacing = 1.0

    matrix_text = styles.add_style("Matrix Text", WD_STYLE_TYPE.PARAGRAPH)
    matrix_text.font.name = "Calibri"
    matrix_text.font.size = Pt(7.4)
    matrix_text.paragraph_format.space_after = Pt(1)
    matrix_text.paragraph_format.line_spacing = 0.95

    quote = styles.add_style("Key Finding", WD_STYLE_TYPE.PARAGRAPH)
    quote.font.name = "Calibri"
    quote.font.size = Pt(11)
    quote.font.bold = True
    quote.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    quote.paragraph_format.space_after = Pt(0)
    quote.paragraph_format.line_spacing = 1.15

    ref = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    ref.font.name = "Calibri"
    ref.font.size = Pt(8.5)
    ref.paragraph_format.left_indent = Inches(0.25)
    ref.paragraph_format.first_line_indent = Inches(-0.25)
    ref.paragraph_format.space_after = Pt(4)
    ref.paragraph_format.line_spacing = 1.0


def configure_headers_footers(doc: Document) -> None:
    for section in doc.sections:
        hp = section.header.paragraphs[0]
        hp.text = "PATIENT-SPECIFIC ECG CONDUCTION–REPOLARIZATION REVIEW"
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in hp.runs:
            run.font.name = "Calibri"
            run.font.size = Pt(7.5)
            run.font.color.rgb = RGBColor.from_string("68737D")
        fp = section.footer.paragraphs[0]
        add_page_number(fp)


def add_para(doc, text="", style=None, bold_lead=None, align=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(4)
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_numbered(doc, items):
    for index, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.28)
        p.paragraph_format.first_line_indent = Inches(-0.28)
        p.paragraph_format.space_after = Pt(4)
        n = p.add_run(f"{index}. ")
        n.bold = True
        if isinstance(item, tuple):
            lead, rest = item
            r = p.add_run(lead)
            r.bold = True
            p.add_run(rest)
        else:
            p.add_run(item)


def add_callout(doc, heading, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_cant_split(table.rows[0])
    set_table_indent(table, 0)
    set_table_borders(table, color=fill, size=8)
    cell = table.cell(0, 0)
    set_cell_width(cell, 9360)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.style = doc.styles["Key Finding"]
    p.add_run(heading + "  ").bold = True
    r = p.add_run(text)
    r.bold = False
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_table(doc, headers, rows, widths, font_style="Table Text", header_fill=LIGHT_GRAY):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_indent(table, 120)
    set_table_borders(table)
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_cant_split(hdr)
    for j, (head, width) in enumerate(zip(headers, widths)):
        cell = hdr.cells[j]
        set_cell_width(cell, width)
        set_cell_shading(cell, header_fill)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.style = doc.styles[font_style]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(head)
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        set_cant_split(table.rows[-1])
        if i % 2:
            for c in cells:
                set_cell_shading(c, "FBFCFD")
        for j, (value, width) in enumerate(zip(row, widths)):
            cell = cells[j]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            p = cell.paragraphs[0]
            p.style = doc.styles[font_style]
            p.add_run(str(value))
    if len(rows) <= 6:
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    set_keep_with_next(paragraph, True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_reference(doc, n, citation, doi=None, pmid=None, url=None):
    p = doc.add_paragraph(style="Reference")
    p.add_run(f"{n}. {citation}")
    links = []
    if doi:
        links.append(("DOI", f"https://doi.org/{doi}"))
    if pmid:
        links.append((f"PMID {pmid}", f"https://pubmed.ncbi.nlm.nih.gov/{pmid}/"))
    if url:
        links.append(("Full text", url))
    if links:
        p.add_run(" ")
        for idx, (label, href) in enumerate(links):
            if idx:
                p.add_run(" · ")
            add_hyperlink(p, label, href)


def add_title_page(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("CLINICAL ECG METHODS REVIEW")
    r.font.size = Pt(9)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)

    p = doc.add_paragraph(style="Title")
    p.add_run("Patient-Specific ECG Conduction and Repolarization Changes for Seizure Detection")

    p = doc.add_paragraph(style="Report Subtitle")
    p.add_run("Physiological evidence, signal-processing methods, artifact constraints, and a defensible third-branch research design")

    meta = doc.add_table(rows=3, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.LEFT
    meta.autofit = False
    set_table_indent(meta, 0)
    set_table_borders(meta, color=WHITE, size=0)
    vals = [
        ("Prepared for", "ECG seizure-detection feature architecture"),
        ("Review date", "19 August 2026"),
        ("Review type", "Rapid structured critical review; not a registered systematic review"),
    ]
    for i, (a, b) in enumerate(vals):
        set_cell_width(meta.cell(i, 0), 1800)
        set_cell_width(meta.cell(i, 1), 7560)
        set_cell_margins(meta.cell(i, 0), start=0)
        set_cell_margins(meta.cell(i, 1), start=0)
        p1 = meta.cell(i, 0).paragraphs[0]
        p1.style = doc.styles["Small Body"]
        r1 = p1.add_run(a.upper())
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(BLUE)
        p2 = meta.cell(i, 1).paragraphs[0]
        p2.style = doc.styles["Small Body"]
        p2.add_run(b)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    add_callout(
        doc,
        "Bottom line",
        "A new branch is scientifically defensible if it measures patient-specific, heart-rate-conditioned deviations in conduction and repolarization and exposes uncertainty. It is not defensible as a generic ‘ST abnormality detector,’ because absolute ST voltage is filter-, posture-, lead-, ectopy-, and motion-sensitive and has almost no direct seizure-detection validation.",
        PALE_GREEN,
    )
    p = doc.add_paragraph(style="Small Body")
    p.paragraph_format.space_before = Pt(14)
    p.add_run("Clinical caution. ").bold = True
    p.add_run("This document proposes research features and validation tests. It is not a diagnostic rule, a substitute for clinical ECG interpretation, or evidence that a detected anomaly is caused by a seizure.")
    doc.add_page_break()


def build_document():
    doc = Document()
    set_section_geometry(doc.sections[0])
    configure_styles(doc)
    configure_headers_footers(doc)
    add_title_page(doc)

    doc.add_heading("1. Executive summary", level=1)
    add_callout(
        doc,
        "Recommended branch",
        "Patient-Calibrated Conduction–Repolarization Anomaly (PCRA): an eligibility-gated branch that models PR–heart-rate, QRS-duration, QT/JT–RR-history, and ST–T morphology residuals relative to each patient’s own stable baseline, then summarizes persistent change over windows.",
    )
    add_para(doc, "The professor’s proposed ‘conduction branch’ is best interpreted as a third, physiologically organized view of the ECG—not as another copy of the RR branch and not as a universal arrhythmia classifier. Its research question is: during an analyzable sequence of sinus-conducted beats, does atrioventricular conduction, ventricular depolarization, or ventricular repolarization depart from what is expected for this patient at the current rate and recent RR history?")
    add_bullets(doc, [
        ("What direct seizure evidence supports. ", "Peri-ictal PR behavior can change; QT/QTc can lengthen or shorten; postictal T-wave alternans can rise after generalized convulsive seizures; and rare bradyarrhythmia, AV block, or asystole can occur. Direction and prevalence are heterogeneous across patients and seizure types [3–11]."),
        ("What it does not support. ", "There is no validated, universal seizure-specific ST-segment voltage pattern, JT signature, or single conduction threshold. Evidence for patient-specific ST anomaly detection during seizures is essentially absent."),
        ("Why personalization matters. ", "QT–RR adaptation and hysteresis are strongly subject-specific, and seizure ECG detector performance is concentrated in physiological ‘responders.’ A patient baseline is therefore a principled necessity, not just a machine-learning preference [20,24]."),
        ("Why frequency/power methods have not solved it. ", "The best-established repolarization spectrum is T-wave alternans power at exactly 0.5 cycles/beat. Generic FFT or wavelet power of raw ST–T samples has no seizure-specific physiological label and is readily dominated by rate, QRS misalignment, muscle activity, motion, and filtering [28–30]."),
        ("What should be shipped as research output. ", "Separate scores for eligibility, PR–HR residual change, QRS change, QT/JT–RR-history residual change, ST–T morphology change, T-wave alternans, and rare safety events—plus confidence. Do not initially collapse them into one unexplained score."),
    ])

    doc.add_heading("Decision for the current project", level=2)
    add_table(doc,
        ["Decision", "Recommendation", "Reason"],
        [
            ("Create the branch?", "Yes—as an exploratory, interpretable branch.", "The physiology and transferable methods justify a testable branch, but not a clinical claim."),
            ("Make ST deviation its center?", "No.", "Absolute ST amplitude is too measurement- and context-sensitive and has little direct seizure evidence."),
            ("Primary modeling idea", "Conditional patient residuals + persistence.", "Removes predictable rate/history effects and reduces isolated-noise triggers."),
            ("PVCs/ectopy", "Gate or model as a separate beat class.", "PVCs change depolarization and produce secondary, usually discordant ST–T changes; mixing them with sinus beats creates false ‘repolarization’ anomalies."),
            ("Relationship to existing branches", "Late fusion with RR and morphology branches.", "The new branch should prove incremental value rather than let a classifier rediscover heart-rate or generic waveform artifacts."),
        ],
        [1500, 3000, 4860],
    )

    doc.add_heading("2. Scope, terminology, and the measurement problem", level=1)
    doc.add_heading("2.1 Three different biological questions", level=2)
    add_table(doc,
        ["Domain", "ECG observables", "Question answered", "Main confounders"],
        [
            ("Conduction", "P duration; PR interval; PR segment; AV block pattern", "How atrial activation reaches the ventricles", "P-wave visibility; AV nodal rate dependence; ectopy; drugs"),
            ("Ventricular depolarization", "QRS duration and morphology", "How ventricular activation propagates", "PVCs; bundle-branch block; pacing; lead change"),
            ("Ventricular repolarization", "J point; ST level/slope/area; T morphology; QT, JT, Tpeak–Tend; TWA", "How ventricular recovery evolves", "Heart rate/history; secondary changes from abnormal QRS; posture; electrolytes; ischemia; filters"),
            ("Eligibility", "Signal quality, beat class, landmark confidence", "Whether a physiological measurement is trustworthy", "Motion, EMG, electrode contact, clipping, baseline drift"),
        ],
        [1500, 2400, 2700, 2760],
    )
    add_para(doc, "The branch name may retain ‘conduction’ for architectural consistency, but the scientific label should explicitly include repolarization. QT and ST–T features are not conduction features in the narrow electrophysiological sense. Keeping the concepts distinct prevents an uninterpretable feature bucket.")

    doc.add_heading("2.2 True ST-interval shortening versus ST-segment voltage deviation", level=2)
    add_callout(
        doc,
        "These are not the same change",
        "ST duration is a time measurement from the J point to T-wave onset (when T onset can be delineated). ST deviation is a voltage measurement relative to an isoelectric reference, usually at the J point or a defined delay after J. A shorter duration suggests altered timing of early repolarization; elevation/depression reflects a shifted potential. Neither can be inferred from the other [1].",
        PALE_AMBER,
    )
    add_para(doc, "In a single wearable lead, T-wave onset is often less reproducible than T-wave end. The practical implementation should therefore avoid calling J-to-T-onset a robust ‘ST interval’ unless landmark confidence is demonstrated. For amplitude, the measurement definition must include the reference baseline, evaluation point, units, lead, and filter chain. A phrase such as ‘ST changed’ is scientifically incomplete.")

    doc.add_heading("2.3 Primary and secondary repolarization changes", level=2)
    add_para(doc, "Primary repolarization change arises from altered recovery despite broadly normal activation. Secondary repolarization change follows altered ventricular activation: a wide or abnormal QRS changes the expected ST–T direction and shape [1]. A PVC is the clearest example. It usually has a premature, wide QRS without a normally conducted preceding P wave, followed by secondary ST–T discordance. The branch must therefore classify the beat before interpreting its ST–T segment.")
    add_bullets(doc, [
        "Sinus-conducted beats: eligible for patient residuals when landmarks and signal quality are acceptable.",
        "PVCs, paced beats, fusion beats, aberrantly conducted beats: exclude from sinus repolarization statistics or maintain a separate ectopy track.",
        "Persistent bundle-branch block or pacing: establish a separate patient baseline; do not compare it with narrow-QRS baseline beats.",
        "New sustained QRS widening or AV block: emit a safety/arrhythmia event, not a seizure label.",
    ])

    doc.add_heading("3. Physiological rationale", level=1)
    add_para(doc, "Seizures can alter autonomic output, catecholamine release, respiration, oxygenation, acid–base status, and cortical–cardiac network activity. These mechanisms can change sinus rate, AV nodal conduction, repolarization adaptation, and arrhythmia susceptibility. The ECG is therefore a plausible peripheral proxy, but it is not specific: exercise, arousal, pain, sleep transitions, medications, hypoxemia, electrolyte changes, and movement can produce overlapping patterns.")
    add_table(doc,
        ["Mechanism", "Expected ECG pathway", "Feature implication", "Specificity problem"],
        [
            ("Sympathetic surge", "Tachycardia; shorter AV nodal delay; altered QT adaptation", "PR–HR and QT–RR residuals, slopes, persistence", "Exercise, stress, awakening"),
            ("Parasympathetic activation", "Bradycardia, AV delay/block, pauses", "PR prolongation, dropped beats, safety flags", "Sleep, medications, vasovagal events"),
            ("Hypoxemia / respiratory dysfunction", "Repolarization instability and QTc abnormality", "QT/JT residuals; T morphology; oxygen covariate", "Apnea and pulmonary disease"),
            ("Abnormal ventricular activation", "Wide QRS with secondary ST–T change", "Beat-type-aware QRS/ST–T interpretation", "PVCs, BBB, pacing"),
            ("Transient repolarization lability", "Beat-to-beat QT/T-wave alternation", "STVQT, QTVI, TWA", "Noise, ectopy, nonstationarity"),
        ],
        [1700, 2600, 2600, 2460],
    )
    add_para(doc, "The correct inferential target is consequently not ‘did the ECG become abnormal?’ but ‘did a high-quality, physiologically coherent deviation occur beyond the patient’s expected response to rate, history, state, and beat type?’")

    doc.add_heading("4. Review methods and evidence grading", level=1)
    add_para(doc, "This was a rapid structured critical review designed to support feature-branch architecture. It was not prospectively registered and should not be presented as a PRISMA-compliant systematic review. Searches were run on 19 August 2026. PubMed/MEDLINE, NCBI full text, Crossref/publisher pages, ClinicalTrials.gov, and public IEEE/Google Scholar discoverability were used. Embase, Scopus, and Web of Science were not directly searched because authenticated subscriptions were unavailable; this is a coverage limitation.")
    doc.add_heading("4.1 Search strategy and counts", level=2)
    add_para(doc, "The anchor PubMed query combined seizure/ictal terms, ECG terms, and conduction–repolarization terms in title/abstract fields and returned 575 records. Four additional focused PubMed searches returned 174 feature-focused, 178 detection-focused, 50 personalization/anomaly, and 347 conduction-safety records; overlap was substantial and counts are not additive. ClinicalTrials.gov returned 34 broad records, of which approximately ten were potentially relevant by title and none clearly targeted patient-specific ST/conduction residuals. Crossref and publisher searches were used for DOI verification rather than quantitative screening because broad Crossref counts were non-specific.")
    add_para(doc, "A pre-existing project corpus had catalogued 47 unique full-text PDFs. The current review reused its traceable extraction, rechecked central full texts and abstracts, and citation-chained forward and backward from the key seizure ECG studies. Thirty-three core publications are summarized in the evidence matrix. Because the broad result set was not independently dual-screened, no precise PRISMA exclusion flow is claimed.")
    add_table(doc,
        ["Class", "Meaning", "How it is used here"],
        [
            ("A", "Direct peri-ictal conduction/repolarization evidence", "Supports biological plausibility, timing, direction, or safety events."),
            ("B", "Direct ECG seizure detection/prediction/anomaly study", "Supports windowing, personalization, evaluation, or achievable performance—not necessarily a specific conduction feature."),
            ("C", "Epilepsy-associated interictal/postictal or disease-group physiology", "Supports vulnerability or biomarker plausibility but not event detection."),
            ("D", "Transferable ECG signal-processing method outside seizure detection", "Supports measurement or anomaly technique only."),
            ("E", "Standards, consensus, or general algorithm reference", "Defines correct measurement and interpretation."),
        ],
        [850, 3500, 5010],
    )

    doc.add_heading("4.2 Eligibility criteria", level=2)
    add_bullets(doc, [
        ("Included direct evidence: ", "human ECG with seizure timing or epilepsy status and an explicit PR, QRS, QT/QTc, JT, ST, T-wave, TWA, block, or full-complex endpoint; or ECG-only seizure detection with patient-level evaluation."),
        ("Included transferable evidence: ", "validated delineation, QT–RR hysteresis, QT variability, TWA, ST monitoring, robust anomaly detection, or change-point methods that can be mapped to a single-lead wearable workflow."),
        ("Excluded from core claims: ", "animal-only physiology, EEG-only methods, HRV-only work when it offered no transferable anomaly/evaluation principle, case reports as prevalence evidence, and studies without analyzable ECG methods."),
    ])

    doc.add_heading("5. Human seizure evidence", level=1)
    doc.add_heading("5.1 PR interval and atrioventricular conduction", level=2)
    add_para(doc, "Pensel et al. analyzed 56 mesial temporal lobe seizures in 14 patients. Average PR shortened during seizures, but individual direction varied; the PR–heart-rate relationship differed by seizure lateralization. No PR interval exceeded 200 ms and no complete AV block occurred in that series [3]. The finding supports modeling the conditional PR–HR relationship rather than a fixed PR threshold. It does not establish a stand-alone detector.")
    add_para(doc, "Older series and case aggregation show that clinically important AV block or asystole can occur, but these events are rare and heterogeneous [10,11]. They should be handled as safety outputs with clinician review. Their rarity makes them unsuitable as the main seizure feature.")

    doc.add_heading("5.2 QRS duration and ventricular depolarization", level=2)
    add_para(doc, "Direct seizure studies rarely report beatwise QRS-duration dynamics after excluding ectopy and movement. Diab et al. used the full P–Q–R–S–T complex in a supervised detector, demonstrating that full-complex information may contribute [16], but the paper does not validate a seizure-specific QRS widening mechanism. QRS duration is therefore a useful control and secondary feature: it detects a change in activation and helps decide whether any ST–T change is primary or secondary.")

    doc.add_heading("5.3 QT/QTc: real signal, contradictory direction", level=2)
    add_para(doc, "The QT literature is the strongest direct repolarization evidence and also the clearest warning against universal thresholds. Nei et al. reported peri-ictal rhythm/repolarization abnormalities in 39% of 51 seizures but no significant mean PR or QTc shift [4]. Brotherstone et al. found QTc lengthening in 156 seizures from 39 patients, with formula-dependent abnormality [5]. Surges et al. instead found enhanced QTc shortening, particularly early after secondarily generalized tonic–clonic seizures [6]. Seyal et al. linked seizure-related QTc abnormalities to ictal hypoxemia [7], while Moseley et al. documented ECG and oximetry changes in partial-complex and generalized seizures [8].")
    add_callout(doc, "Interpretation", "Both QT lengthening and shortening may be meaningful. A signed population rule is unlikely to generalize. Model the magnitude and direction of each patient’s QT/JT residual after accounting for effective RR and hysteresis, and preserve the sign for later phenotype analysis.", PALE_AMBER)
    add_para(doc, "Bazett correction alone is inadequate because it overcorrects at high heart rates, exactly where many seizures operate. Report raw QT, JT, RR, and several correction models for comparison, but make a patient-specific QT–RR model with explicit rate history the primary representation [24,25].")

    doc.add_heading("5.4 ST segment, T-wave morphology, and alternans", level=2)
    add_para(doc, "There is no convincing event-level literature establishing patient-specific ST elevation/depression as a reliable seizure detector. Some older studies grouped ‘ST–T’ or repolarization abnormalities without modern signal-quality and filter controls [4]. This is insufficient for a core ST-voltage feature.")
    add_para(doc, "T-wave alternans has more specific evidence. Strzelczyk et al. found an increase after generalized tonic–clonic seizures in 16 patients [9]. Pang et al. found higher ambulatory TWA in chronic than newly diagnosed epilepsy, which supports repolarization instability but not seizure timing [14]. A 2026 pediatric surgery study found interictal repolarization abnormalities improved after seizure resolution, again supporting disease-state biology rather than an event detector [15]. These studies justify a carefully gated TWA sub-analysis, especially postictally, not a universal real-time trigger.")

    doc.add_heading("5.5 Rare arrhythmias are safety events, not branch labels", level=2)
    add_para(doc, "The multicenter peri-ictal arrhythmia study by Vilella et al. analyzed 455 generalized convulsive seizures in 249 patients and found tachycardia to be common while clinically important bradyarrhythmias and conduction events were uncommon [10]. Its explicit artifact exclusions also show why tonic–clonic epochs are the hardest place to claim subtle ST–T morphology. The branch should emit pauses, high-grade AV block, asystole, new sustained wide-complex rhythm, and extreme QT as separate urgent events. They must not be silently fused into a probabilistic seizure score.")

    doc.add_heading("6. Measurement and signal-processing methods", level=1)
    doc.add_heading("6.1 Landmark and beat chain", level=2)
    add_numbered(doc, [
        ("Create separate signal streams. ", "Use a QRS-detection stream optimized for robust R peaks and a morphology stream that preserves diagnostic low-frequency content. Never assume one aggressive bandpass is suitable for both."),
        ("Delineate with confidence. ", "Estimate P onset/offset, QRS onset/offset, J point, T peak, T onset when reliable, and T end. Store each landmark’s confidence and local quality."),
        ("Classify beats before measurement. ", "Sinus, supraventricular ectopic, PVC, paced, fusion, noise/unknown. Create homogeneous beat chains."),
        ("Measure without premature correction. ", "Store raw PR, QRS, QT, JT, Tpeak–Tend, ST values, RR, and context. Derived corrections remain reproducible only if raw quantities are retained."),
        ("Condition on history. ", "Compute effective RR from a causal adaptation kernel or patient-specific hysteresis model instead of using the immediately preceding RR alone."),
    ])
    add_para(doc, "Automatic delineation is feasible—modern convolutional approaches show good performance on mixed-quality annotated ECG sets [31]—but algorithm-level precision does not guarantee that beat-to-beat variation is physiologically resolved. The project’s existing prototype correctly treats endpoint error as a validation gate: if P95 landmark error approaches or exceeds expected physiological variability, variability features cannot be trusted.")

    doc.add_heading("6.2 Per-beat feature definitions", level=2)
    add_table(doc,
        ["Feature", "Definition", "Patient-conditioned form", "Status"],
        [
            ("PR", "P onset to QRS onset", "PR − expected PR(HR, state)", "Secondary/direct evidence"),
            ("PR segment", "P offset to QRS onset", "Residual conditional on HR", "Exploratory; P offset fragile"),
            ("QRSd", "QRS onset to QRS offset", "Residual within beat-type/template class", "Core control / secondary feature"),
            ("QT", "QRS onset to T end", "QT − expected QT(effective RR, patient)", "Core repolarization feature"),
            ("JT", "QRS offset/J point to T end", "JT residual; compare with QT residual", "Secondary; less direct seizure evidence"),
            ("Tpeak–Tend", "T peak to T end", "Patient-template residual", "Secondary/exploratory"),
            ("ST level", "Voltage relative to TP/PR reference at J and fixed/rate-adjusted offsets", "Within-lead baseline residual", "Exploratory only"),
            ("ST slope/area", "Trend or integral over defined J-to-T window", "Template/context residual", "Exploratory; sensitive to drift"),
            ("T morphology", "Amplitude, area, asymmetry, width, template/DTW/PCA distance", "Distance to patient’s rate/context template", "Secondary/exploratory"),
        ],
        [1300, 2700, 3100, 2260],
    )

    doc.add_heading("6.3 QT and repolarization variability", level=2)
    add_para(doc, "Three constructs must remain separate:")
    add_bullets(doc, [
        ("QT variability (QTV): ", "beat-to-beat fluctuation in measured QT after landmark and artifact control."),
        ("QT variability index (QTVI): ", "log10 of normalized QT variance divided by normalized heart-rate variance, historically evaluated over relatively stable multi-minute windows [26]. It is not a short-window seizure score."),
        ("Short-term QT variability (STVQT): ", "the average orthogonal distance between successive QT values, commonly computed over 30 beats: Σ|QTn+1 − QTn|/(N√2) [27]. It remains highly sensitive to T-end jitter."),
    ])
    add_para(doc, "Consensus guidance emphasizes that normal stable beat-to-beat QTV is small, often on the order of only a few milliseconds, so T-end measurement quality is decisive [25]. The development pipeline should compare observed variability with repeated-annotation and algorithm-error distributions; otherwise it risks measuring delineator noise.")

    doc.add_heading("6.4 T-wave alternans, morphology variability, and frequency/power", level=2)
    add_table(doc,
        ["Method", "What it measures", "Typical requirement", "Noise behavior", "Seizure status"],
        [
            ("Spectral TWA", "Alternating T-wave amplitude/shape power at 0.5 cycles/beat versus noise floor", "Stable, aligned sinus-beat sequence; often long windows", "Sensitive to ectopy, nonstationarity, and alignment", "Direct small postictal study [9]; otherwise transferable [28]"),
            ("Modified moving average", "Difference between recursively updated odd/even beat templates", "Alternating A/B beat streams; artifact rejection", "More adaptable to ambulatory data; still motion-sensitive", "Epilepsy disease-group evidence [14]; method validated elsewhere [29,30]"),
            ("Template distance / DTW", "Shape change after alignment or nonlinear warping", "Rate- and beat-class-specific reference templates", "Can absorb benign timing variation; may also hide true change", "Transferable, not seizure-validated"),
            ("PCA morphology", "Energy outside dominant patient T-wave subspace", "Consistent lead and enough clean templates", "Sensitive to lead rotation and T-end errors", "Transferable, usually multilead"),
            ("Generic FFT/wavelet energy", "Broad ST–T frequency content", "Explicit band/window definition", "Easily dominated by QRS leakage, rate, EMG, motion", "No convincing seizure-specific validation"),
        ],
        [1600, 2700, 1900, 1900, 1260],
    )
    add_callout(doc, "Answer to ‘should there be a power-level method?’", "Yes, but only with a physiological target. TWA power at 0.5 cycles/beat is interpretable. HRV LF/HF power belongs in the RR branch. Broad raw ST–T spectral power is an exploratory morphology descriptor, not established evidence of conduction or seizure activity.", PALE_GREEN)

    doc.add_heading("7. Patient-specific anomaly and window dynamics", level=1)
    doc.add_heading("7.1 Why conditional baselines are necessary", level=2)
    add_para(doc, "A patient’s expected QT depends on current and preceding RR intervals; QT adaptation is delayed, nonlinear, and stable enough to be subject-specific [24]. PR also depends on heart rate and autonomic state. ST–T morphology depends on lead vector, posture, respiration, and activation sequence. A single global median is therefore not an adequate baseline.")
    add_para(doc, "For beat i, define a feature residual as observed minus expected under the patient model. Examples are rQT,i = QTi − E[QT | effective RRi, patient, lead/state] and rPR,i = PRi − E[PR | HRi, patient, state]. For morphology, use a distance to the nearest valid patient template within the same beat class and rate/context stratum.")

    doc.add_heading("7.2 Calibration set", level=2)
    add_bullets(doc, [
        "Use EEG-confirmed seizure-free data when available; otherwise exclude wide safety margins around reported/detected events.",
        "Sample multiple sleep/wake periods, postures, activity levels, and heart-rate ranges. One hand-picked ‘clean’ 4-minute interval is too narrow for deployment.",
        "Require sufficient support in each stratum; for development, target hundreds of eligible sinus beats per rate/context cluster and at least several independent time blocks.",
        "Freeze the baseline around alarms, seizures, acute illness, medication changes, and lead reapplication. Online adaptation should occur only during high-confidence, low-anomaly periods and be audited for drift.",
        "Maintain separate baselines when persistent QRS morphology changes indicate a new lead placement, bundle-branch pattern, or pacing state.",
    ])

    doc.add_heading("7.3 Robust anomaly estimators", level=2)
    add_table(doc,
        ["Estimator", "Use in branch", "Advantage", "Main caution"],
        [
            ("Median/MAD score", "Per-feature robust standardized residual", "Transparent; small calibration burden", "Univariate; MAD collapses in overly stable or quantized signals"),
            ("Robust covariance / MCD", "Joint residual distance across features", "Models correlations; resistant to contamination", "Needs enough baseline samples relative to dimension [36]"),
            ("One-class SVM", "Nonlinear normal-region boundary", "Flexible multivariate shape", "Kernel/tuning opaque; calibration-sensitive [35]"),
            ("LOF", "Local density abnormality", "Can detect context-specific deviations", "Density unstable in high dimension; baseline coverage matters [34]"),
            ("Matrix Profile", "Discord of waveform subsequences", "No hand-crafted landmarks", "May capture rate/envelope/artifact rather than physiology [21,37]"),
            ("CUSUM / change point", "Persistent mean shift in residual stream", "Causal and interpretable", "Requires drift/noise calibration; multiple testing [38,39]"),
        ],
        [1600, 2600, 2500, 2660],
    )
    add_para(doc, "The recommended initial benchmark is median/MAD plus robust covariance. It is auditable and reveals whether added complexity is warranted. OCSVM, LOF, and learned embeddings should be prespecified comparisons, not the only implementation.")

    doc.add_heading("7.4 Window features and persistence", level=2)
    add_para(doc, "Compute window features only over eligible homogeneous beats and report the denominator. Useful summaries are median signed residual, median absolute residual, upper quantile, MAD/IQR, fraction beyond a patient threshold, longest consecutive run, robust Theil–Sen slope, positive and negative CUSUM, and joint robust distance. A window with fewer than the required valid beats should output missing/low confidence, not zero anomaly.")
    add_table(doc,
        ["Time scale", "Purpose", "Candidate outputs"],
        [
            ("Beat / 3–10 beats", "Abrupt conduction or activation event", "AV block, pause, QRS widening, PVC burst; safety and eligibility"),
            ("20–60 beats", "Ictal transition and persistence", "Residual median, run length, slope, CUSUM, template distance"),
            ("~30 eligible beats", "Short-term QT variability", "STVQT with uncertainty and T-end quality"),
            ("≥128 stable beats", "TWA / longer variability", "Spectral TWA or MMA; QTVI only if stationarity is adequate"),
            ("Minutes to hours", "Baseline drift and context", "Rate/history model, circadian/activity strata, lead-state change"),
        ],
        [1700, 3500, 4160],
    )

    doc.add_heading("8. Noise, filtering, and artifact robustness", level=1)
    add_callout(doc, "Hard constraint", "If the acquisition or preprocessing has removed the low-frequency content required for ST measurement, no downstream model can reconstruct trustworthy absolute ST deviation. A 0.5–40 Hz stream can be acceptable for R peaks or coarse anomaly detection but must not be presented as diagnostic-band ST measurement [2,32].", PALE_RED)
    doc.add_heading("8.1 Filter architecture", level=2)
    add_bullets(doc, [
        ("Morphology stream: ", "retain raw data and use a diagnostic-preserving response; AHA guidance supports a low-frequency cutoff near 0.05 Hz for ST fidelity. Offline zero-phase processing and online causal processing must be validated separately [2]."),
        ("Detection stream: ", "a narrower band can improve R-peak detection, but its output is not used for ST level or T morphology."),
        ("Do not normalize away amplitude: ", "per-window z-normalization can help generic anomaly algorithms but destroys absolute ST/T amplitude meaning."),
        ("Validate the entire chain: ", "inject synthetic and recorded ST steps, slow ramps, T-wave changes, and QRS shifts through acquisition plus filtering; quantify bias, settling, ringing, and latency."),
    ])
    add_para(doc, "High-pass filters at 0.5 Hz can create or suppress apparent ST shifts, especially with causal implementations and large QRS/T complexes [2,32]. The SeizeIT2 raw-waveform anomaly benchmark filtered 0.5–40 Hz and downsampled to 8 Hz [21]. It therefore supports coarse patient ECG anomaly detection, not fiducial timing or diagnostic ST voltage.")

    doc.add_heading("8.2 Artifact sources and required gates", level=2)
    add_table(doc,
        ["Artifact/confounder", "Failure mode", "Required defense"],
        [
            ("Baseline wander / respiration", "False ST shift, area, and T-wave asymmetry", "Baseline estimator excluding QRS/ST–T; diagnostic-band validation; respiratory/context covariate"),
            ("EMG and convulsive movement", "T-end jitter, false TWA, waveform discord", "High-frequency noise metric, accelerometer, template residual, multi-beat persistence, reject tonic artifact"),
            ("Electrode contact/pop or lead rotation", "Abrupt persistent ST/T template change", "Saturation/step detector; lead-state clustering; freeze/recalibrate baseline"),
            ("Clipping/saturation/dropout", "Artificial plateaus and interval errors", "Hard invalidation and missing-data accounting"),
            ("PVCs/ectopy", "Wide QRS and secondary ST–T discordance", "Beat classification; separate ectopy stream; exclude adjacent contaminated beats"),
            ("Heart-rate acceleration", "Predictable PR/QT/T change", "Conditional residuals using current rate and RR history"),
            ("Posture/activity", "Lead-vector and ST/T morphology shift", "Activity/posture strata; hard negative test sets"),
            ("Drugs/electrolytes/ischemia", "True non-seizure repolarization change", "Clinical covariates; do not claim seizure specificity; safety review"),
        ],
        [2050, 3300, 4010],
    )

    doc.add_heading("8.3 Artifact-centered validation", level=2)
    add_numbered(doc, [
        "Evaluate clean seizures, artifact-heavy seizures, and analyzable seizure subintervals separately.",
        "Create hard negatives matched for heart rate and movement: exercise, stairs, awakening, coughing, sleep transitions, posture changes, electrode disturbances, and ectopy runs.",
        "Report performance stratified by signal-quality decile and fraction of time eligible. A detector that abstains on every difficult seizure is not robust.",
        "Perform blinded beat-level review of a sample of high anomaly windows by an ECG-competent annotator, labeling physiology versus artifact and primary versus secondary ST–T change.",
        "Measure residual correlation with accelerometer magnitude, baseline wander, high-frequency noise, and QRS-template change. Strong correlation indicates artifact leakage.",
        "Run negative-control features placed in physiologically irrelevant intervals; if they perform similarly, the model is probably exploiting motion or timing artifacts.",
    ])

    doc.add_heading("9. Existing ECG seizure-detection approaches", level=1)
    add_para(doc, "Most ECG seizure detectors are dominated by rate and HRV. They validate patient personalization and event-level evaluation, but they do not prove that conduction/repolarization adds information. That incremental question must be tested explicitly.")
    add_table(doc,
        ["Study", "Data and method", "Reported result", "What transfers", "What does not"],
        [
            ("Fujiwara 2016 [17]", "14 patients; multivariate statistical process control on 8 HRV features", "91% sensitivity; ~0.7 false alarms/hour in selected awake episodes", "Patient baseline; Hotelling-like anomaly statistics", "Conduction/repolarization mechanism"),
            ("Billeci 2018 [18]", "Patient-specific HRV and recurrence quantification", "Promising patient-specific prediction", "Personalization and nonlinear dynamics", "Direct ST/QT evidence"),
            ("Karasmanoglou 2023 [19]", "5 women, 10 focal seizures; LOF/MCD/OCSVM on 2–3 min HRV windows", "LOF AUC ~97%, sensitivity ~93%, specificity ~96% under selected reference labels", "Normal-only anomaly framework", "Generalization; movement robustness; conduction"),
            ("Jeppesen 2025 [20]", "Prospective phase 3 personalized HR threshold; selected >50 bpm responders", "42 seizures/17 eligible; 90.5% sensitivity; 2.5 alarms/day", "Prospective personalization and responder concept", "Unselected population performance; repolarization"),
            ("Diab 2025 [16]", "32 patients, 47 seizures; PQRST-derived beat features; Extra Trees", "60-beat/20-trigger: 86% sensitivity, 99.9% specificity; ~1.5 false alarms/hour", "Full-complex feasibility; persistence", "External validation; diagnostic ST fidelity; patient baseline"),
            ("Reintjes 2025 [21]", "SeizeIT2, 856 usable seizures/120 patients; raw 0.5–40 Hz ECG downsampled to 8 Hz; three anomaly methods", "Strong sensitivity–false-alarm tradeoff; responder dependence", "Large open patient-level benchmark", "Fiducial/ST interpretation"),
            ("Alhaskir 2026 [22]", "236 patients; wearable ECG HRV; leave-one-patient-out", "Median sensitivity 66.6%; median 5.2 false alarms/24 h", "Large focal-seizure reality check", "Conduction/repolarization-specific value"),
        ],
        [1300, 3000, 2100, 1900, 1060],
    )
    add_para(doc, "A critical evaluation detail is event definition. Reintjes et al. showed that widening the accepted detection window from strict ictal overlap to −5/+3 minutes changes both sensitivity and false-alarm interpretation [21]. The new branch should report strict onset-centered results and clinically motivated extended windows separately.")

    doc.add_heading("10. Evidence matrix", level=1)
    add_para(doc, "The matrix below grades what each source can support. ‘Direct’ does not mean clinically validated; many direct studies are small, inpatient, single-lead, or artifact-limited.")
    matrix_rows = [
        ("A", "Pensel 2021 [3]", "14 / 56 mesial temporal seizures", "Single-lead VEEG", "PR and RR; PR–HR relation", "Mean PR shortening; individual directions vary; no high-grade block", "Small, inpatient; lateralization exploratory"),
        ("A", "Nei 2000 [4]", "43 / 51 refractory partial seizures", "VEEG ECG", "Rhythm/repolarization; PR/QTc", "39% any abnormality; no significant mean PR/QTc shift", "Older methods; heterogeneous labels"),
        ("A", "Brotherstone 2010 [5]", "39 / 156 seizures", "Inpatient ECG", "QTc, four correction formulas", "Significant QTc lengthening; formula-dependent abnormality", "9-beat averages; no deployment test"),
        ("A", "Surges 2010 [6]", "25 patients", "VEEG ECG", "QTc and postictal HR", "Enhanced QT shortening; greater after generalized seizures", "Small; correction/hysteresis concerns"),
        ("A", "Seyal 2011 [7]", "Refractory epilepsy seizures", "VEEG ECG + oximetry", "QTc and hypoxemia", "Repolarization abnormalities associated with ictal hypoxemia", "Association, not causal detector"),
        ("A", "Moseley 2011 [8]", "Partial complex/generalized seizures", "ECG + oximetry", "ECG and oxygen changes", "Documents peri-ictal cardiac/oximetric change", "Small; mixed seizure types"),
        ("A", "Strzelczyk 2011 [9]", "16 patients", "VEEG ECG", "Postictal TWA", "TWA increased after generalized tonic–clonic seizures", "Small; postictal; algorithm details require full text"),
        ("A", "Vilella 2024 [10]", "249 / 455 generalized convulsive seizures", "Prospective multicenter VEEG", "Peri-ictal arrhythmias", "Tachycardia common; serious conduction events rare", "Ictal movement limits subtle morphology"),
        ("A", "Gigli 2023 [13]", "117 patients", "Basal vs postictal 12-lead", "QTc, Brugada/ERP, ECG abnormalities", "44% postictal abnormal; many abnormalities also basal", "Postictal association; no detector"),
        ("A", "van der Lende 2016 [11]", "Systematic review/case aggregation", "Published peri-ictal ECG", "Bradyarrhythmia, AV block, asystole", "Rare but clinically important events characterized", "Publication bias; case-level denominators"),
        ("B", "Diab 2025 [16]", "32 / 47 seizures", "Thoracic ECG, 256/512 Hz", "PQRST beat features; Extra Trees; 3–60 beats", "86% sensitivity at selected 60-beat setting; 1.5 FA/h", "Selected clean ECG; internal validation"),
        ("B", "Fujiwara 2016 [17]", "14 patients; 11 awake preictal episodes", "Long-term ECG", "8 HRV features; MSPC", "91% sensitivity; ~0.7 FA/h", "Small selected set; HRV only"),
        ("B", "Billeci 2018 [18]", "Patient-specific seizure set", "ECG", "HRV + recurrence quantification", "Supports patient-specific nonlinear prediction", "Not conduction/repolarization"),
        ("B", "Karasmanoglou 2023 [19]", "5 women / 10 focal seizures", "200 Hz ECG", "LOF, MCD, OCSVM on HRV", "High AUC under selected reference labeling", "Manual stable baseline; tiny sample"),
        ("B", "Jeppesen 2025 [20]", "101 enrolled; 42 seizures/17 eligible", "Wearable ECG + smartphone", "Personalized HR threshold", "90.5% sensitivity; 2.5 alarms/day", "Responder-enriched eligibility"),
        ("B", "Reintjes 2025 [21]", "120 / 856 usable seizures", "Single-lead wearable ECG", "Matrix Profile, MADRID, TimeVQVAE-AD", "Open large benchmark; substantial FA tradeoff", "8 Hz after 0.5–40 Hz; coarse signal"),
        ("B", "Alhaskir 2026 [22]", "236 / 260 seizures", "Wearable ECG", "HRV; LOPO validation", "Median 66.6% sensitivity; 5.2 FA/24h", "HRV only; variability across patients"),
        ("B", "Bhagubai 2025 [23]", "125 / 883 seizures; >11,000 h", "Wearable EEG + ECG", "Open multimodal dataset", "Enables patient-level waveform testing", "Dataset paper, not feature validation"),
        ("C", "Ali 2017 [12]", "59 children after convulsive SE + 31 controls", "12-lead/clinical ECG", "ST/T, QTc, QT–RR variability", "Epilepsy associated with post-SE ventricular alterations", "Status epilepticus; group comparison"),
        ("C", "Pang 2019 [14]", "6 newly diagnosed vs 6 chronic", "Holter/patch", "Modified moving-average TWA", "Higher TWA in chronic epilepsy", "Interictal disease marker; very small"),
        ("C", "Bozdag 2026 [15]", "Pediatric drug-resistant epilepsy", "Interictal ECG", "Repolarization before/after surgery", "Abnormalities improve after seizure resolution", "Not event detection; pediatric surgical cohort"),
        ("D", "Malik 2008 [24]", "40 subjects × 3 long recordings", "Long ECG", "Subject-specific QT/RR hysteresis", "Profiles individual and stable", "Non-seizure physiology"),
        ("D/E", "Baumert 2016 [25]", "Consensus", "Surface ECG", "QTV definitions and quality", "Measurement guidance; normal QTV small", "No seizure validation"),
        ("D", "Berger 1997 [26]", "Cardiomyopathy vs controls", "256-s ECG", "QTVI", "Introduces normalized QT/HR variability index", "Long stable window; disease context"),
        ("D", "Hinterseer 2008 [27]", "Drug-induced long-QT pilot", "30-beat ECG", "STVQT", "Increased short-term QT variability", "T-end error-sensitive; no seizures"),
        ("D/E", "Verrier 2011 [28]", "Consensus", "Holter/exercise ECG", "Spectral and time-domain TWA", "Defines 0.5 cycles/beat physiology and quality", "No seizure detector validation"),
        ("D", "Nearing 2002 [29]", "Experimental validation", "ECG", "Modified moving-average TWA", "Robust alternating-template approach", "Non-seizure endpoint"),
        ("D", "Selvaraj 2009 [30]", "Ambulatory ECG simulations/data", "Holter ECG", "MMA vs spectral TWA under noise", "Quantifies method-specific noise behavior", "No seizures"),
        ("D", "Jimenez-Perez 2021 [31]", "QTDB, 105 two-lead records", "Annotated ECG", "CNN delineation", "Supports automated fiducial extraction", "Dataset differs from wearable seizure ECG"),
        ("D/E", "Buendía-Fuentes 2012 [32]", "Filter analysis", "Clinical ECG", "High-pass effect on ST", "0.5 Hz filters can create ST interpretation errors", "Not ambulatory seizure data"),
        ("D", "Shusterman 2007 [33]", "Ischemia monitoring", "Surface ECG", "Dynamic ST tracking", "Transferable baseline and persistence concepts", "Different pathology and leads"),
        ("E", "AHA ST/QT 2009 [1]", "Scientific statement", "Standard ECG", "ST, T, U, QT definitions", "Measurement and primary/secondary interpretation", "Not a detector study"),
        ("E", "AHA technology 2007 [2]", "Scientific statement", "Acquisition/filters", "Bandwidth and distortion", "Defines morphology-preserving constraints", "Not wearable-specific"),
    ]
    add_table(doc,
        ["Class", "Source", "Population", "Setting", "Endpoint/method", "Main result", "Critical limitation"],
        matrix_rows,
        [540, 1300, 1550, 1150, 1700, 1900, 1220],
        font_style="Matrix Text",
    )

    doc.add_heading("11. Critical gaps and feasibility", level=1)
    add_table(doc,
        ["Gap", "Current evidence", "Consequence for the branch"],
        [
            ("Universal direction", "QT lengthening and shortening are both reported; PR direction varies.", "Use signed patient residuals and phenotype analysis, not one-sided rules."),
            ("ST-segment seizure signature", "No robust patient-specific event-detection evidence.", "ST level/slope/area remain exploratory and require strict filter provenance."),
            ("JT and Tpeak–Tend", "Plausible, but little direct seizure validation.", "Secondary features with explicit missingness/confidence."),
            ("T-wave morphology", "Disease-state and transferable methods exceed event-level seizure evidence.", "Template-distance experiments, not claims of established biomarkers."),
            ("Artifact-rich convulsions", "Subtle fiducials often fail during the movements of greatest clinical interest.", "Eligibility coverage is a primary endpoint; multimodal quality indicators required."),
            ("Patient-specific calibration", "Strong rationale, few conduction/repolarization seizure implementations.", "This is a research contribution and must be validated prospectively."),
            ("Incremental value", "Most published ECG performance can be explained by rate/HRV.", "Ablate RR and morphology branches; report added value and redundancy."),
            ("Generalization", "Many studies are small, internal, or responder-enriched.", "Patient-level external validation and calibration reporting are mandatory."),
        ],
        [2200, 3500, 3660],
    )
    add_para(doc, "The idea is feasible because all proposed core variables can be extracted from a sufficiently sampled single-lead ECG. The scientific risk is not lack of possible features; it is measurement validity and causal ambiguity. A model can learn a seizure-correlated pattern while actually using tachycardia, motion, lead shift, or convulsive EMG. The design must be built to falsify those shortcuts.")

    doc.add_heading("12. Proposed research design: PCRA branch", level=1)
    add_callout(doc, "Design principle", "Treat eligibility as a modeled state, patient baseline as conditional physiology, and persistence as evidence. Preserve separate domain outputs until incremental value and artifact independence are demonstrated.", PALE_GREEN)

    doc.add_heading("12.1 Branch architecture", level=2)
    pipeline_rows = [
        ("0. Provenance", "Sampling rate, units, lead placement, raw availability, filter impulse/phase response", "Reject unsupported diagnostic claims"),
        ("1. Eligibility", "Signal quality, saturation, baseline wander, EMG, contact step, beat type, landmark confidence", "Beat/domain mask + confidence"),
        ("2. Per-beat physiology", "PR, QRSd, QT, JT, Tpeak–Tend, defined ST measures, T template coordinates", "Raw values with uncertainties"),
        ("3. Patient expectation", "Robust PR–HR model; QT–effective-RR hysteresis; rate/context template banks", "Expected value and residual"),
        ("4. Window dynamics", "Median, quantiles, MAD, run length, slope, CUSUM, STVQT/QTVI/TWA when eligible", "Persistent domain evidence"),
        ("5. Outputs", "Aeligible, APR, AQRS, AQT/JT, AST–T, ATWA, Asafety, uncertainty", "Separate calibrated scores"),
        ("6. Fusion", "Late fusion with RR and morphology branches; missingness-aware model", "Final seizure probability only after ablations"),
    ]
    add_table(doc, ["Stage", "Computation", "Output"], pipeline_rows, [1550, 5100, 2710], header_fill=PALE_BLUE)

    doc.add_heading("12.2 Core preprocessing specification", level=2)
    add_bullets(doc, [
        "Retain an immutable raw stream. Record ADC units, gain, sampling rate, lead geometry, and clock continuity.",
        "Use morphology-preserving filtering for P/ST/T analysis; validate causal and offline versions with known waveforms. Maintain a separate QRS-detection stream.",
        "Estimate signal quality per domain: P eligibility can fail while QRS/QT remains valid; TWA may require a stricter mask than QT.",
        "Detect and classify ectopy before computing sinus-chain features. Exclude at least the ectopic beat and physiologically contaminated neighbors from QT/TWA chains according to a prespecified rule.",
        "Resample or align beats only with documented interpolation error. Never infer TWA from alternation created by sub-sample misalignment.",
    ])

    doc.add_heading("12.3 Patient baseline and adaptation", level=2)
    add_table(doc,
        ["Component", "Initial calibration", "Deployment adaptation", "Failure safeguard"],
        [
            ("PR–HR", "Robust regression/splines over eligible sinus beats and rate strata", "Slow update during high-quality low-anomaly periods", "Freeze near alarms and medication/lead changes"),
            ("QT/JT–RR", "Patient-specific nonlinear QT–RR plus causal hysteresis kernel", "Update intercept slowly; re-estimate curve only with adequate range", "Do not update from tachycardic seizure/postictal epochs"),
            ("QRS/T templates", "Clusters by beat class, lead state, rate/context", "Add template only after repeated stable confirmation", "Novel persistent cluster triggers lead-state review"),
            ("ST baseline", "Within-lead, posture/activity-aware distribution", "Conservative; preferably session-specific", "Disable absolute interpretation after lead disturbance"),
            ("Thresholds", "Patient quantiles with minimum sample support", "Monitor false-alarm drift; clinician-controlled reset", "Never adapt to suppress confirmed seizures"),
        ],
        [1550, 3200, 2800, 1810],
    )

    doc.add_heading("12.4 Output contract", level=2)
    add_para(doc, "At each analysis time, the branch should expose a structured feature vector rather than a single opaque anomaly:")
    add_bullets(doc, [
        ("Eligibility: ", "valid-beat fraction, domain-specific quality, landmark uncertainty, beat-class composition, lead-state confidence."),
        ("PR domain: ", "signed and absolute PR–HR residual summaries, slope, persistence, block/pauses."),
        ("QRS domain: ", "duration residual, template distance, new-wide-complex flag."),
        ("QT/JT domain: ", "raw intervals, patient QT/JT–effective-RR residuals, STVQT, optional QTVI, uncertainty."),
        ("ST–T domain: ", "defined ST level/slope/area residuals and T-template distance, with filter/lead validity flag."),
        ("TWA domain: ", "MMA and/or 0.5-cycles/beat spectral estimate, noise floor, eligible-beat count, ectopy correction metadata."),
        ("Safety: ", "high-grade AV block, asystole/pause, sustained brady/tachyarrhythmia, extreme QT; independent clinical escalation path."),
    ])

    doc.add_heading("12.5 Evaluation plan", level=2)
    add_numbered(doc, [
        ("Data split. ", "Patient-level train/validation/test split or leave-one-patient-out; nested tuning. No beat/window leakage across the same patient or seizure."),
        ("Reference labels. ", "Video-EEG onset/offset with seizure type, state, medication, oxygen, and artifact annotations where available."),
        ("Primary metrics. ", "Event sensitivity, false alarms per 24 hours, median detection latency, time-in-warning, PPV, and analyzable-time coverage with patient-clustered 95% intervals."),
        ("Secondary metrics. ", "Per-patient calibration, responder fraction, feature availability, quality-stratified performance, and false-alarm taxonomy."),
        ("Comparator models. ", "RR branch alone; morphology branch alone; RR+morphology; PCRA alone; all branches; generic raw-waveform anomaly baseline."),
        ("Hard-negative cohorts. ", "Rate-matched activity, sleep transitions, posture/lead changes, ectopy, and non-seizure hypoxemia."),
        ("Prospective freeze. ", "After retrospective development, lock definitions, preprocessing, thresholds, and fusion before prospective validation."),
    ])

    doc.add_heading("12.6 Prespecified ablations", level=2)
    add_table(doc,
        ["Ablation", "Question answered"],
        [
            ("Remove HR/RR and replace residuals with raw intervals", "Does patient rate/history correction add information or merely transform heart rate?"),
            ("Remove eligibility/persistence", "How much apparent performance comes from noisy isolated beats?"),
            ("Exclude all ST/T amplitude features", "Is interval timing alone sufficient and more robust?"),
            ("Exclude QT; retain JT", "Is the signal driven by depolarization onset/QRS effects?"),
            ("Exclude ectopy and post-ectopic beats versus model separately", "Are results caused by PVC-associated secondary repolarization?"),
            ("Generic population baseline versus patient baseline", "Does personalization improve calibration and false alarms?"),
            ("Strict ictal window versus extended peri-ictal window", "Is performance truly detection, or broad temporal association?"),
            ("Diagnostic-preserving stream versus 0.5–40 Hz stream", "How much morphology signal is filter-dependent or artifactual?"),
        ],
        [3900, 5460],
    )

    doc.add_heading("13. Testable hypotheses and feature ranking", level=1)
    doc.add_heading("13.1 Falsifiable hypotheses", level=2)
    add_numbered(doc, [
        "Patient-specific QT/JT residual magnitude will separate peri-ictal from matched interictal windows better than Bazett QTc or raw QT alone.",
        "A PR–HR residual score will add modest but significant event-level information beyond RR features in a subset of patients, without a universal signed direction.",
        "Eligibility gating and persistence will reduce false alarms from motion/ectopy more than they reduce seizure sensitivity.",
        "T-wave template distance will add value primarily in high-quality non-convulsive or postictal intervals, not during tonic–clonic motion.",
        "TWA will rise postictally after generalized convulsive seizures in a reproducible subgroup but will not be a high-coverage early detector.",
        "Absolute ST voltage will lose most apparent association after filter-provenance, posture/activity, rate, and lead-state matching. This is an important expected-null test.",
        "The full PCRA branch will improve false alarms per 24 hours or patient-level calibration beyond RR+morphology, even if overall AUROC changes little.",
        "Residual patterns will form patient phenotypes—QT-shortening, QT-lengthening, PR-dominant, TWA-dominant, or non-responder—rather than a single population response.",
    ])

    doc.add_heading("13.2 Feature ranking", level=2)
    add_table(doc,
        ["Tier", "Features", "Rationale / condition"],
        [
            ("Core infrastructure", "Raw provenance; domain eligibility; beat class; landmark uncertainty; valid-beat fraction", "Without these, interval and ST–T claims are not interpretable."),
            ("Core physiology", "PR–HR residual; QRS-duration residual/control; QT and JT residuals using effective RR/hysteresis; signed/absolute robust window summaries", "Best balance of plausibility, interpretability, and single-lead feasibility."),
            ("Secondary", "Tpeak–Tend; T-template distance; STVQT; PR/QT slopes and CUSUM", "Potential incremental value; needs stronger precision and stability evidence."),
            ("Exploratory", "Defined ST level/slope/area residuals; PCA/DTW; QTVI; spectral TWA; MMA TWA; one-class models; change points", "Methodologically defensible but limited seizure validation or stricter data requirements."),
            ("Safety-only", "High-grade AV block, asystole/pause, sustained arrhythmia, new wide-complex pattern, extreme QT", "Clinically important, too rare/non-specific to act as seizure classifiers."),
            ("Do not prioritize", "Generic raw ST–T power; Bazett-only QTc; mixed sinus/PVC morphology; random window splits; scores without eligibility", "High artifact/redundancy risk and weak physiological interpretability."),
        ],
        [1550, 3900, 3910],
    )

    doc.add_heading("13.3 Go/no-go criteria", level=2)
    add_bullets(doc, [
        ("Measurement go: ", "landmark error and repeatability are materially smaller than the patient changes being modeled, with adequate analyzable coverage."),
        ("Artifact go: ", "anomaly is not reproduced by matched motion, posture, ectopy, or filter perturbation controls."),
        ("Incremental go: ", "the branch improves patient-level false-alarm burden, sensitivity, latency, or calibration beyond RR+morphology under locked evaluation."),
        ("Clinical go: ", "external prospective validation confirms performance and safety-event handling. Until then, label the branch exploratory."),
    ])
    add_callout(doc, "Recommended immediate next experiment", "Implement only the core infrastructure and core physiology tiers first. On a small, carefully adjudicated patient set, plot PR–HR and QT/JT–effective-RR residual trajectories around seizures alongside quality, beat type, RR-branch output, and accelerometry. If the residuals disappear when artifacts and rate history are controlled, stop before adding dozens of spectral features.", PALE_GREEN)

    doc.add_heading("References", level=1)
    refs = [
        ("Rautaharju PM, et al. AHA/ACCF/HRS recommendations for standardization and interpretation of the electrocardiogram: Part IV—the ST segment, T and U waves, and the QT interval. J Am Coll Cardiol. 2009;53:982–991.", "10.1016/j.jacc.2008.12.014", "19281931", None),
        ("Kligfield P, et al. Recommendations for standardization and interpretation of the electrocardiogram: Part I—the electrocardiogram and its technology. J Am Coll Cardiol. 2007;49:1109–1127.", "10.1016/j.jacc.2007.01.024", "17349896", None),
        ("Pensel MC, Basili LM, Jordan A, Surges R. Atrioventricular conduction in mesial temporal lobe seizures. Front Neurol. 2021;12:661391.", "10.3389/fneur.2021.661391", "33995256", "https://pmc.ncbi.nlm.nih.gov/articles/PMC8115552/"),
        ("Nei M, Ho RT, Sperling MR. EKG abnormalities during partial seizures in refractory epilepsy. Epilepsia. 2000;41:542–548.", "10.1111/j.1528-1157.2000.tb00207.x", "10802759", None),
        ("Brotherstone R, Blackhall B, McLellan A. Lengthening of corrected QT during epileptic seizures. Epilepsia. 2010;51:221–232.", "10.1111/j.1528-1167.2009.02281.x", "19732135", None),
        ("Surges R, Scott CA, Walker MC. Enhanced QT shortening and persistent tachycardia after generalized seizures. Neurology. 2010;74:421–426.", "10.1212/WNL.0b013e3181ccc706", "20124208", "https://pmc.ncbi.nlm.nih.gov/articles/PMC2872619/"),
        ("Seyal M, Pascual F, Lee CY, Li CS, Bateman LM. Seizure-related cardiac repolarization abnormalities are associated with ictal hypoxemia. Epilepsia. 2011;52:2105–2111.", "10.1111/j.1528-1167.2011.03262.x", "21906052", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3203996/"),
        ("Moseley BD, Wirrell EC, Nickels K, Johnson JN, Ackerman MJ, Britton J. Electrocardiographic and oximetric changes during partial complex and generalized seizures. Epilepsy Res. 2011;95:237–245.", "10.1016/j.eplepsyres.2011.04.005", "21561737", None),
        ("Strzelczyk A, et al. Postictal increase in T-wave alternans after generalized tonic-clonic seizures. Epilepsia. 2011;52:2112–2117.", "10.1111/j.1528-1167.2011.03266.x", "21933179", None),
        ("Vilella L, et al. Incidence and types of cardiac arrhythmias in the peri-ictal period in patients having a generalized convulsive seizure. Neurology. 2024;103:e209501.", "10.1212/WNL.0000000000209501", "38870452", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11759939/"),
        ("van der Lende M, Surges R, Sander JW, Thijs RD. Cardiac arrhythmias during or after epileptic seizures. J Neurol Neurosurg Psychiatry. 2016;87:69–74.", "10.1136/jnnp-2015-310559", "26038597", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4717443/"),
        ("Ali W, et al. Epilepsy is associated with ventricular alterations following convulsive status epilepticus in children. Epilepsia Open. 2017;2:432–440.", "10.1002/epi4.12074", "29430560", "https://pmc.ncbi.nlm.nih.gov/articles/PMC5800777/"),
        ("Gigli L, et al. Electrocardiogram changes in the postictal phase of epileptic seizure: results from a prospective study. J Clin Med. 2023;12:4098.", "10.3390/jcm12124098", "37373791", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10299131/"),
        ("Pang TD, et al. Cardiac electrical instability in newly diagnosed/chronic epilepsy tracked by Holter and ECG patch. Neurology. 2019;93:450–458.", "10.1212/WNL.0000000000008077", "31477610", None),
        ("Bozdag E, et al. Interictal cardiac repolarization abnormalities improve after surgical seizure resolution in pediatric drug-resistant epilepsy. Epilepsia. 2026;67:3034–3047.", "10.1002/epi.70206", "41848583", None),
        ("Diab E, et al. Electrocardiogram (ECG)-based seizure detection using supervised machine-learning. Neurophysiol Clin. 2025;55:103098.", "10.1016/j.neucli.2025.103098", None, None),
        ("Fujiwara K, et al. Epileptic seizure prediction based on multivariate statistical process control of heart rate variability features. IEEE Trans Biomed Eng. 2016;63:1321–1332.", "10.1109/TBME.2015.2512276", "26841385", None),
        ("Billeci L, Marino D, Insana L, Vatti G, Varanini M. Patient-specific seizure prediction based on heart rate variability and recurrence quantification analysis. PLoS One. 2018;13:e0204339.", "10.1371/journal.pone.0204339", "30252915", "https://pmc.ncbi.nlm.nih.gov/articles/PMC6155519/"),
        ("Karasmanoglou A, Antonakakis M, Zervakis M. ECG-based semi-supervised anomaly detection for early detection and monitoring of epileptic seizures. Int J Environ Res Public Health. 2023;20:5000.", "10.3390/ijerph20065000", "36981911", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10049350/"),
        ("Jeppesen J, et al. Seizure detection using wearable electrocardiogram connected to a smartphone: a phase 3 clinical validation study. EBioMedicine. 2025;120:105952.", "10.1016/j.ebiom.2025.105952", "41027311", "https://pmc.ncbi.nlm.nih.gov/articles/PMC12516532/"),
        ("Reintjes C, Hagenbeck JF, Ballo M, Rahlmeier T, Wolf SM, Schoder D. ECG-based detection of epileptic seizures in real-world wearable settings: insights from the SeizeIT2 dataset. Sensors. 2025;25:7687.", "10.3390/s25247687", "41471682", "https://pmc.ncbi.nlm.nih.gov/articles/PMC12736484/"),
        ("Alhaskir M, et al. Reliable detection of focal onset impaired awareness seizures in patients with epilepsy using wearable ECG: development and validation study. Comput Methods Programs Biomed. 2026;279:109295.", "10.1016/j.cmpb.2026.109295", "41764784", None),
        ("Bhagubai M, et al. SeizeIT2: wearable dataset of patients with focal epilepsy. Sci Data. 2025;12:1228.", "10.1038/s41597-025-05580-x", "40664714", None),
        ("Malik M, Hnatkova K, Novotny T, Schmidt G. Subject-specific profiles of QT/RR hysteresis. Am J Physiol Heart Circ Physiol. 2008;295:H2356–H2363.", "10.1152/ajpheart.00625.2008", "18849333", None),
        ("Baumert M, et al. QT interval variability in body surface ECG: measurement, physiological basis, and clinical value—EHRA/ESC consensus guidance. Europace. 2016;18:925–944.", "10.1093/europace/euv405", "26823389", None),
        ("Berger RD, et al. Beat-to-beat QT interval variability: novel evidence for repolarization lability in ischemic and nonischemic dilated cardiomyopathy. Circulation. 1997;96:1557–1565.", "10.1161/01.CIR.96.5.1557", "9315547", None),
        ("Hinterseer M, et al. Beat-to-beat variability of QT intervals is increased in patients with drug-induced long-QT syndrome: a case-control pilot study. Eur Heart J. 2008;29:185–190.", "10.1093/eurheartj/ehm586", "18156612", None),
        ("Verrier RL, et al. Microvolt T-wave alternans: physiological basis, methods of measurement, and clinical utility—consensus guideline. J Am Coll Cardiol. 2011;58:1309–1324.", "10.1016/j.jacc.2011.06.029", "21920259", None),
        ("Nearing BD, Verrier RL. Modified moving average analysis of T-wave alternans to predict ventricular fibrillation with high accuracy. J Appl Physiol. 2002;92:541–549.", "10.1152/japplphysiol.00592.2001", "11796662", None),
        ("Selvaraj RJ, Chauhan VS. Effect of noise on T-wave alternans measurement in ambulatory ECGs using modified moving average versus spectral method. Pacing Clin Electrophysiol. 2009;32:632–641.", "10.1111/j.1540-8159.2009.02337.x", "19422585", None),
        ("Jimenez-Perez G, Alcaine A, Camara O. Delineation of the electrocardiogram with a mixed-quality-annotations dataset using convolutional neural networks. Sci Rep. 2021;11:863.", "10.1038/s41598-020-79512-7", "33441632", None),
        ("Buendía-Fuentes F, et al. High-bandpass filters in electrocardiography: source of error in the interpretation of the ST segment. ISRN Cardiol. 2012;2012:706217.", "10.5402/2012/706217", "22778996", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3388307/"),
        ("Shusterman V, Goldberg A, Schindler DM, Fleischmann KE, Lux RL, Drew BJ. Dynamic tracking of ischemia in the surface electrocardiogram. J Electrocardiol. 2007;40:S179–S186.", "10.1016/j.jelectrocard.2007.06.015", "17993319", None),
        ("Breunig MM, Kriegel H-P, Ng RT, Sander J. LOF: identifying density-based local outliers. Proc ACM SIGMOD. 2000:93–104.", "10.1145/335191.335388", None, None),
        ("Schölkopf B, Platt JC, Shawe-Taylor J, Smola AJ, Williamson RC. Estimating the support of a high-dimensional distribution. Neural Comput. 2001;13:1443–1471.", "10.1162/089976601750264965", None, None),
        ("Rousseeuw PJ, Van Driessen K. A fast algorithm for the minimum covariance determinant estimator. Technometrics. 1999;41:212–223.", "10.1080/00401706.1999.10485670", None, None),
        ("Yeh C-CM, Zhu Y, Ulanova L, et al. Matrix Profile I: all pairs similarity joins for time series. IEEE ICDM. 2016:1317–1322.", "10.1109/ICDM.2016.0179", None, None),
        ("Page ES. Continuous inspection schemes. Biometrika. 1954;41:100–115.", "10.1093/biomet/41.1-2.100", None, None),
        ("Killick R, Fearnhead P, Eckley IA. Optimal detection of changepoints with a linear computational cost. J Am Stat Assoc. 2012;107:1590–1598.", "10.1080/01621459.2012.737745", None, None),
        ("Bukhari HA, et al. Estimation of potassium levels in hemodialysis patients by T wave nonlinear dynamics and morphology markers. Comput Biol Med. 2022;143:105304.", "10.1016/j.compbiomed.2022.105304", "35168084", None),
    ]
    for i, (citation, doi, pmid, url) in enumerate(refs, 1):
        add_reference(doc, i, citation, doi=doi, pmid=pmid, url=url)

    doc.add_heading("Appendix A. Reproducible search strings", level=1)
    searches = [
        ("PubMed anchor", "((epilepsy[Title/Abstract] OR seizure*[Title/Abstract] OR ictal[Title/Abstract] OR peri-ictal[Title/Abstract] OR postictal[Title/Abstract]) AND (electrocardiogram[Title/Abstract] OR ECG[Title/Abstract] OR EKG[Title/Abstract] OR Holter[Title/Abstract]) AND (\"PR interval\"[Title/Abstract] OR \"atrioventricular conduction\"[Title/Abstract] OR \"QRS duration\"[Title/Abstract] OR \"QT interval\"[Title/Abstract] OR QTc[Title/Abstract] OR \"JT interval\"[Title/Abstract] OR \"ST segment\"[Title/Abstract] OR \"T wave\"[Title/Abstract] OR \"T-wave alternans\"[Title/Abstract] OR repolarization[Title/Abstract] OR \"AV block\"[Title/Abstract] OR asystole[Title/Abstract]))", "575"),
        ("Feature-focused", "(seizure*[Title] OR ictal[Title] OR postictal[Title] OR epilepsy[Title]) AND (ECG[Title/Abstract] OR electrocardiogra*[Title/Abstract]) AND (\"PR interval\" OR \"QRS duration\" OR \"QT interval\" OR QTc OR \"ST segment\" OR \"T-wave\" OR repolarization)", "174"),
        ("Detection-focused", "(\"seizure detection\"[Title/Abstract] OR \"seizure prediction\"[Title/Abstract]) AND (ECG[Title/Abstract] OR electrocardiogra*[Title/Abstract] OR cardiac[Title/Abstract])", "178"),
        ("Personalization", "(seizure* OR epilepsy) AND (ECG OR electrocardiogra* OR heart rate variability) AND (patient-specific OR personalized OR anomaly detection OR one-class)", "50"),
        ("Conduction safety", "(seizure*[Title] OR ictal[Title] OR epilepsy[Title]) AND (atrioventricular block OR AV block OR asystole OR bradyarrhythmia)", "347"),
        ("ClinicalTrials.gov", "(epilepsy OR seizure) AND (ECG OR electrocardiogram) AND detection", "34 broad; ~10 potentially relevant by title"),
    ]
    add_table(doc, ["Source", "Search string", "Results on 2026-08-19"], searches, [1500, 6300, 1560], font_style="Matrix Text")
    add_para(doc, "Search notes: PubMed query counts are reproducible snapshots and will change as records are indexed. The focused search counts overlap and must not be summed. Public web search and citation chaining were used to identify publisher pages, standards, and post-cutoff papers. Lack of direct Embase/Scopus/Web of Science access and absence of dual screening limit claims of exhaustiveness.", style="Small Body")

    doc.add_heading("Appendix B. Minimum branch data dictionary", level=1)
    add_table(doc,
        ["Field group", "Required fields"],
        [
            ("Provenance", "patient_id; recording_id; lead; sampling_hz; gain/units; device; filter chain; raw_available; clock status"),
            ("Beat identity", "beat_id/time; RR; beat class; R confidence; preceding/following ectopy; lead-state cluster"),
            ("Landmarks", "P onset/offset; QRS onset/offset; J; T onset/peak/end; confidence and uncertainty for each"),
            ("Raw features", "PR; PR segment; QRSd; QT; JT; Tpeak–Tend; ST reference/locations; ST slope/area; T morphology coordinates"),
            ("Expected/residual", "patient model version; context stratum; effective RR; expected value; signed residual; robust standardized residual"),
            ("Quality", "baseline wander; high-frequency noise; template correlation; saturation/dropout; motion; domain eligibility mask"),
            ("Window", "start/end; valid beats/total; residual summaries; run length; slope/CUSUM; STVQT/QTVI/TWA; uncertainty"),
            ("Labels", "EEG seizure onset/offset/type; sleep/activity; artifact; oxygen; medication/clinical context; review status"),
        ],
        [2100, 7260],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
